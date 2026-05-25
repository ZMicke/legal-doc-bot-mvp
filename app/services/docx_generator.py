import re
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.config import settings


SECTION_PATTERN = re.compile(r"^\d+\.\s+")


def clean_markdown(text: str) -> str:
    text = text.replace("**", "")
    text = text.replace("__", "")
    text = re.sub(r"^\s*[-*]\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    return text.strip()


def apply_run_style(run, size: int = 12, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold


def add_paragraph(
    document: Document,
    text: str = "",
    *,
    bold: bool = False,
    alignment: int | None = None,
    size: int = 12,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)

    if alignment is not None:
        paragraph.alignment = alignment

    run = paragraph.add_run(text)
    apply_run_style(run, size=size, bold=bold)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(0.8)

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.15


def generate_claim_docx(claim_text: str) -> str:
    generated_dir = Path(settings.GENERATED_DIR)
    generated_dir.mkdir(parents=True, exist_ok=True)

    file_id = str(uuid4())
    file_path = generated_dir / f"claim_{file_id}.docx"

    document = Document()
    configure_document(document)

    clean_text = clean_markdown(claim_text)

    for line in clean_text.splitlines():
        clean_line = line.strip()

        if not clean_line:
            document.add_paragraph()
            continue

        if clean_line == "ПРЕТЕНЗИЯ":
            add_paragraph(
                document,
                clean_line,
                bold=True,
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
                size=14,
            )
            continue

        if clean_line == "о нарушении срока оказания услуг":
            add_paragraph(
                document,
                clean_line,
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
            )
            continue

        if SECTION_PATTERN.match(clean_line) and not clean_line.endswith("."):
            add_paragraph(document, clean_line, bold=True)
            continue

        add_paragraph(document, clean_line)

    document.save(file_path)

    return str(file_path)
